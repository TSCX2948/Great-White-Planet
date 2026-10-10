"""文档文本提取：把一个文件变成 AI 能读的内容。

扩展方式（需求文档第十二节：未来支持 txt/ppt/pptx/docx/xls/xlsx）
--------------------------------------------------------------
新增一种文件类型时，只需在 :data:`READERS` 里加一个 ``扩展名 -> 读取函数`` 的映射，
或者用 :func:`register_reader` 注册。其它模块无需改动。

准确性优先原则（需求文档 PDF 处理）
----------------------------------
* PDF 有文本层 → 提取文字；
* 每页文字过少（扫描件）或整体不可靠 → 不猜内容，改为把页面渲染成图片插入 Markdown，
  由用户或支持视觉的 AI 来看图。
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from gwp.core import constants as C

from gwp.services.omml import omml_to_latex

log = logging.getLogger(__name__)


@dataclass
class ExtractedDocument:
    """提取结果。``markdown`` 是可以直接塞给 AI 或写进文档的内容。"""

    path: Path
    markdown: str = ""
    kind: str = "text"                 # text | pdf | docx | pptx | xlsx | image | binary | unknown
    reliable: bool = True              # False 表示"不敢保证识别准确"
    notes: list[str] = field(default_factory=list)
    image_files: list[Path] = field(default_factory=list)
    truncated: bool = False
    error: str = ""

    @property
    def name(self) -> str:
        return self.path.name

    @property
    def is_empty(self) -> bool:
        return not self.markdown.strip()


# --------------------------------------------------------------------------
# 读取器注册表
# --------------------------------------------------------------------------

ReaderFn = Callable[[Path, "ExtractOptions"], ExtractedDocument]

READERS: dict[str, ReaderFn] = {}


def register_reader(extensions: tuple[str, ...] | list[str], reader: ReaderFn) -> None:
    for ext in extensions:
        READERS[ext.lower()] = reader


@dataclass
class ExtractOptions:
    """提取参数，由 config.json 的 summary / import 段填充。"""

    max_chars: int = 200_000
    min_text_chars_per_page: int = 80
    render_pdf_images: bool = True
    image_dir: Path | None = None
    image_max_width: int = 1600
    image_max_pages: int = 60
    keep_table_html: bool = False

    @classmethod
    def from_config(cls, config=None, image_dir: Path | None = None) -> "ExtractOptions":
        if config is None:
            from gwp.core.config import Config

            config = Config.load()
        return cls(
            max_chars=int(config.get("summary.max_input_chars", 200_000) or 200_000),
            min_text_chars_per_page=int(config.get("summary.min_text_chars_per_page", 80) or 80),
            render_pdf_images=bool(config.get("summary.include_images", True)),
            image_dir=image_dir,
            image_max_width=int(config.get("summary.image_max_width", 1600) or 1600),
            image_max_pages=int(config.get("summary.image_max_pages", 60) or 60),
        )


# --------------------------------------------------------------------------
# 纯文本 / Markdown
# --------------------------------------------------------------------------

def read_text_file(path: Path, options: ExtractOptions) -> ExtractedDocument:
    text = ""
    for encoding in ("utf-8", "utf-8-sig", "gbk", "gb18030", "latin-1"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
        except OSError as exc:
            return ExtractedDocument(path=path, kind="text", error=str(exc), reliable=False)
    if not text:
        return ExtractedDocument(path=path, kind="text", markdown="", reliable=False,
                                 notes=["文件为空或无法解码"])
    truncated = len(text) > options.max_chars
    if truncated:
        text = text[: options.max_chars]
    return ExtractedDocument(
        path=path,
        markdown=text,
        kind="text",
        reliable=True,
        truncated=truncated,
        notes=["内容过长，已截断"] if truncated else [],
    )


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def read_pdf(path: Path, options: ExtractOptions) -> ExtractedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ExtractedDocument(
            path=path,
            kind="pdf",
            reliable=False,
            error="未安装 PyMuPDF，无法解析 PDF",
            notes=["请执行：pip install pymupdf"],
        )

    document = ExtractedDocument(path=path, kind="pdf")
    try:
        with fitz.open(path) as pdf:
            if pdf.needs_pass:
                return ExtractedDocument(
                    path=path, kind="pdf", reliable=False,
                    error="PDF 已加密，需要密码", notes=["请先解除密码保护再导入"],
                )
            page_texts: list[str] = []
            for page in pdf:
                try:
                    page_texts.append(page.get_text("text") or "")
                except Exception:  # noqa: BLE001 - 个别损坏页面不应中断整体
                    page_texts.append("")
    except Exception as exc:  # noqa: BLE001
        return ExtractedDocument(path=path, kind="pdf", reliable=False, error=f"PDF 解析失败：{exc}")

    total_pages = len(page_texts)
    char_count = sum(len(t) for t in page_texts)
    avg_per_page = char_count / total_pages if total_pages else 0
    text_ok = avg_per_page >= max(10, options.min_text_chars_per_page)

    if text_ok:
        parts = [f"### 第 {index + 1} 页\n{t.strip()}" for index, t in enumerate(page_texts) if t.strip()]
        markdown = "\n\n".join(parts)
        if not markdown.strip():
            text_ok = False
        else:
            truncated = len(markdown) > options.max_chars
            if truncated:
                markdown = markdown[: options.max_chars]
            document.markdown = markdown
            document.reliable = True
            document.notes.append(f"文本层可用，提取了 {total_pages} 页文字")
            if truncated:
                document.truncated = True
                document.notes.append("内容过长，已截断")
            return document

    # 文本层不可靠 → 转图片（准确性优先，不猜内容）
    document.reliable = False
    document.notes.append(
        f"每页平均仅 {avg_per_page:.0f} 个字符，判定为扫描件/图片型 PDF，改为插入页面图片"
    )
    if not options.render_pdf_images:
        document.markdown = "> 该 PDF 可能是扫描件，且当前设置不允许转为图片。"
        return document

    images = render_pdf_to_images(path, options)
    document.image_files = images
    if not images:
        document.markdown = "> 该 PDF 无法提取文字，也无法渲染为图片。"
        return document
    lines = [f"### 第 {index + 1} 页\n\n![{path.stem} 第 {index + 1} 页]({image.as_posix()})"
             for index, image in enumerate(images)]
    document.markdown = "\n\n".join(lines)
    document.notes.append(f"已把 {len(images)} 页转换为图片插入")
    return document


def render_pdf_to_images(path: Path, options: ExtractOptions) -> list[Path]:
    """把 PDF 每页渲染为 PNG，返回图片路径列表。"""
    try:
        import fitz
    except ImportError:
        return []
    target_dir = options.image_dir
    if target_dir is None:
        target_dir = path.parent / f"{path.stem}_pages"
    target_dir.mkdir(parents=True, exist_ok=True)

    results: list[Path] = []
    try:
        with fitz.open(path) as pdf:
            limit = min(len(pdf), max(1, options.image_max_pages))
            for index in range(limit):
                page = pdf[index]
                zoom = 2.0
                try:
                    width = page.rect.width * zoom
                    if width > options.image_max_width:
                        zoom = options.image_max_width / page.rect.width
                except Exception:  # noqa: BLE001
                    zoom = 1.5
                matrix = fitz.Matrix(zoom, zoom)
                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                target = target_dir / f"{path.stem}_p{index + 1:03d}.png"
                pixmap.save(target)
                results.append(target)
            if len(pdf) > limit:
                log.info("PDF 页数 %s 超过上限 %s，只渲染了前 %s 页", len(pdf), limit, limit)
    except Exception:  # noqa: BLE001
        log.exception("PDF 渲染图片失败: %s", path)
    return results


def extract_pdf_images(path: Path, options: ExtractOptions) -> list[Path]:
    """提取 PDF 内嵌的原始图片（用于保留教材插图）。"""
    try:
        import fitz
    except ImportError:
        return []
    target_dir = options.image_dir or (path.parent / f"{path.stem}_images")
    target_dir.mkdir(parents=True, exist_ok=True)
    results: list[Path] = []
    try:
        with fitz.open(path) as pdf:
            for page_index in range(len(pdf)):
                for image_index, info in enumerate(pdf[page_index].get_images(full=True)):
                    xref = info[0]
                    try:
                        data = pdf.extract_image(xref)
                    except Exception:  # noqa: BLE001
                        continue
                    ext = data.get("ext", "png")
                    target = target_dir / f"{path.stem}_p{page_index + 1}_{image_index}.{ext}"
                    target.write_bytes(data["image"])
                    results.append(target)
    except Exception:  # noqa: BLE001
        log.exception("提取 PDF 内嵌图片失败: %s", path)
    return results


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _docx_omml_text(element) -> str:
    """从一段 OMML（Word 公式的 XML）里把可见文字取出来。

    Word 的公式不是文本 run，而是 ``<m:oMath>`` 这一套 MathML 方言。
    ``paragraph.text`` **完全看不到它** —— 公式会整段消失
    （这就是"公式丢了"的原因）。

    这里只做最保守的事：递归收集 ``m:t``（公式里的文字），拼成一个
    近似可读的字符串（例如 ``ε > 0``）。**不猜结构** ——
    分式/上下标会被拍平成一行，但至少内容不丢，而且调用方可以据此
    把它交给公式渲染器变成图片（需求："公式部分插入图片"）。
    """
    try:
        texts = element.findall(".//" + _OMML_NS + "t")
        return "".join(node.text or "" for node in texts).strip()
    except Exception:  # noqa: BLE001 - 取不到就当没有公式
        return ""


_OMML_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _docx_images(document, path: Path, options: ExtractOptions) -> list[Path]:
    """把 docx 里嵌入的图片导出成文件，返回路径列表。

    为什么必须做：``paragraph.text`` 只返回文字，
    docx 里的插图**完全不会出现**在提取结果里 ——
    对讲义类文档来说那是内容的一部分（示意图、公式截图）。
    """
    from gwp.services import file_ops

    saved: list[Path] = []
    try:
        parts = list(document.part.related_parts.values())
    except Exception:  # noqa: BLE001
        return saved
    target_dir = Path(options.image_dir) if options.image_dir else (
        Path(path).parent / f"{Path(path).stem}_assets"
    )
    seen: set[str] = set()
    for part in parts:
        content_type = getattr(part, "content_type", "") or ""
        if not content_type.startswith("image/"):
            continue
        partname = str(getattr(part, "partname", ""))
        if partname in seen:
            continue
        seen.add(partname)
        suffix = "." + (content_type.split("/")[-1] or "png").replace("jpeg", "jpg")
        try:
            target_dir.mkdir(parents=True, exist_ok=True)
            blob = part.blob
            if len(blob) > options.max_chars * 200:      # 单张图太大就不搬
                continue
            target = file_ops.unique_path(target_dir / f"{Path(path).stem}_img{len(saved)+1}{suffix}")
            target.write_bytes(blob)
            saved.append(target)
        except Exception:  # noqa: BLE001 - 单张图失败不影响其它
            continue
    return saved


def read_docx(path: Path, options: ExtractOptions) -> ExtractedDocument:
    try:
        import docx  # python-docx
        from docx.oxml.ns import qn
    except ImportError:
        return ExtractedDocument(path=path, kind="docx", reliable=False,
                                 error="未安装 python-docx", notes=["请执行：pip install python-docx"])
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        return ExtractedDocument(path=path, kind="docx", reliable=False,
                                 error=f"解析 docx 失败：{exc}")

    try:
        body = document.element.body
    except Exception:  # noqa: BLE001
        body = None

    lines: list[str] = []
    formula_count = 0

    def _paragraph_line(paragraph) -> str:
        """一段文字 -> Markdown 行；**顺带把公式变成 LaTeX 公式**。"""
        nonlocal formula_count
        text = paragraph.text.strip()
        # 公式：paragraph.text 看不到 OMML，所以单独取出来。
        #
        # 需求："公式部分插入图片" —— 要变成图片就得先变成 LaTeX
        # （渲染器吃 LaTeX，而 Word 存的是 OMML）。
        # 转换失败时**退回纯文字**（见 omml_to_latex 的说明），
        # 至少内容不丢。
        #
        # **必须递归找（``.//``）**：Word 的独立公式包在
        # ``<m:oMathPara><m:oMath>…`` 里，只找直接子元素会**整条漏掉**
        # （实测：公式一个都没提取到）。
        if paragraph._p is not None:
            seen_math: set[int] = set()
            for math in paragraph._p.findall(".//" + _OMML_NS + "oMath"):
                if id(math) in seen_math:
                    continue
                seen_math.add(id(math))
                latex = omml_to_latex(math)
                if latex:
                    formula_count += 1
                    # 用 $...$ 包起来：最终渲染时会被换成公式图片
                    text = f"{text} ${latex}$".strip() if text else f"${latex}$"
                else:
                    plain = _docx_omml_text(math)
                    if plain:
                        text = f"{text} {plain}".strip() if text else plain
        if not text:
            return ""
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            return f"{'#' * min(6, int(level) + 1)} {text}"
        if style.startswith("list"):
            return f"- {text}"
        return text

    def _table_lines(table) -> list[str]:
        rows: list[str] = []
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ").replace("|", "\\|")
                for cell in row.cells
            ]
            rows.append("| " + " | ".join(cells) + " |")
        if not rows:
            return []
        separator = "| " + " | ".join("---" for _ in table.columns) + " |"
        # **表格必须是连续的几行**：中间夹空行就渲染不成表格
        # （原实现把每个元素再用空行 join 一次，表格被拆散了）。
        return [rows[0], separator, *rows[1:]]

    # ---- 按**正文顺序**遍历（段落与表格交错出现）----
    #
    # 原实现是"先所有段落、再所有表格"，于是表格全部跑到文档末尾，
    # 上下文全乱。docx 的 body 里子元素本来就是按顺序排的，照着走即可。
    if body is not None:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in body.iterchildren():
            try:
                if child.tag == qn("w:p"):
                    line = _paragraph_line(Paragraph(child, document))
                    if line:
                        lines.append(line)
                elif child.tag == qn("w:tbl"):
                    lines.extend(_table_lines(Table(child, document)))
                    lines.append("")          # 表格后留一个空行（分隔用）
            except Exception:  # noqa: BLE001 - 单个元素失败不该中断整篇
                continue
    else:
        for paragraph in document.paragraphs:
            line = _paragraph_line(paragraph)
            if line:
                lines.append(line)
        for table in document.tables:
            lines.extend(_table_lines(table))
            lines.append("")

    images = _docx_images(document, path, options)
    # 图片以 Markdown 引用挂在末尾（并按需在正文里被消费）
    if images:
        lines.append("")
        for index, image in enumerate(images, start=1):
            lines.append(f"![图 {index}]({image.as_posix()})")

    markdown = "\n".join(lines)
    truncated = len(markdown) > options.max_chars
    if truncated:
        markdown = markdown[: options.max_chars]
    notes: list[str] = []
    if truncated:
        notes.append("内容过长，已截断")
    if formula_count:
        notes.append(f"含 {formula_count} 处公式（已转为公式图片）")
    if images:
        notes.append(f"含 {len(images)} 张内嵌图片")
    return ExtractedDocument(
        path=path, markdown=markdown, kind="docx", reliable=bool(markdown.strip()),
        truncated=truncated, notes=notes, image_files=images,
    )


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

def read_pptx(path: Path, options: ExtractOptions) -> ExtractedDocument:
    try:
        from pptx import Presentation
    except ImportError:
        return ExtractedDocument(path=path, kind="pptx", reliable=False,
                                 error="未安装 python-pptx", notes=["请执行：pip install python-pptx"])
    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        return ExtractedDocument(path=path, kind="pptx", reliable=False,
                                 error=f"解析 pptx 失败：{exc}")

    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                try:
                    for row in shape.table.rows:
                        lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
                except Exception:  # noqa: BLE001
                    pass
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                lines.append(f"> 备注：{note}")
        if lines:
            parts.append(f"### 第 {index} 页幻灯片\n\n" + "\n\n".join(lines))

    markdown = "\n\n".join(parts)
    truncated = len(markdown) > options.max_chars
    if truncated:
        markdown = markdown[: options.max_chars]
    return ExtractedDocument(
        path=path, markdown=markdown, kind="pptx", reliable=bool(markdown.strip()),
        truncated=truncated,
        notes=["仅提取了文字与表格；图表/图片内容未解析"] + (["内容过长，已截断"] if truncated else []),
    )


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def read_xlsx(path: Path, options: ExtractOptions) -> ExtractedDocument:
    try:
        import openpyxl
    except ImportError:
        return ExtractedDocument(path=path, kind="xlsx", reliable=False,
                                 error="未安装 openpyxl", notes=["请执行：pip install openpyxl"])
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        return ExtractedDocument(path=path, kind="xlsx", reliable=False,
                                 error=f"解析 xlsx 失败：{exc}")

    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                if row is None:
                    continue
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    rows.append(cells)
                if len(rows) >= 300:
                    rows.append(["...（已截断，仅保留前 300 行）"])
                    break
            if not rows:
                continue
            width = max(len(r) for r in rows)
            padded = [r + [""] * (width - len(r)) for r in rows]
            lines = ["| " + " | ".join(padded[0]) + " |",
                     "| " + " | ".join("---" for _ in range(width)) + " |"]
            lines.extend("| " + " | ".join(r) + " |" for r in padded[1:])
            parts.append(f"### 工作表：{sheet.title}\n\n" + "\n".join(lines))
    finally:
        try:
            workbook.close()
        except Exception:  # noqa: BLE001
            pass

    markdown = "\n\n".join(parts)
    truncated = len(markdown) > options.max_chars
    if truncated:
        markdown = markdown[: options.max_chars]
    return ExtractedDocument(
        path=path, markdown=markdown, kind="xlsx", reliable=bool(markdown.strip()),
        truncated=truncated,
        notes=["公式以缓存值呈现；图表/图片内容未解析"],
    )


def read_csv_file(path: Path, options: ExtractOptions) -> ExtractedDocument:
    rows: list[list[str]] = []
    try:
        for encoding in ("utf-8-sig", "gbk", "utf-8", "latin-1"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    rows = [row for row in csv.reader(handle) if any(cell.strip() for cell in row)]
                break
            except UnicodeDecodeError:
                continue
    except OSError as exc:
        return ExtractedDocument(path=path, kind="xlsx", reliable=False, error=str(exc))
    if not rows:
        return ExtractedDocument(path=path, markdown="", kind="xlsx", reliable=False)
    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(padded[0]) + " |",
             "| " + " | ".join("---" for _ in range(width)) + " |"]
    lines.extend("| " + " | ".join(r) + " |" for r in padded[1:])
    return ExtractedDocument(path=path, markdown="\n".join(lines), kind="xlsx", reliable=True)


# --------------------------------------------------------------------------
# 图片
# --------------------------------------------------------------------------

def read_image_file(path: Path, options: ExtractOptions) -> ExtractedDocument:
    """图片没有文本：直接以 Markdown 图片语法插入，交由人/AI 视觉查看。"""
    target = path
    if options.image_dir is not None and options.image_max_width > 0:
        target = _copy_or_resize_image(path, options)
    return ExtractedDocument(
        path=path,
        markdown=f"![{path.stem}]({target.as_posix()})",
        kind="image",
        reliable=False,
        image_files=[target],
        notes=["图片内容未做 OCR，按原图插入"],
    )


def _copy_or_resize_image(path: Path, options: ExtractOptions) -> Path:
    """把图片复制到资产目录；超过宽度上限时等比缩小。"""
    target_dir = options.image_dir
    assert target_dir is not None
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / path.name
    try:
        import shutil

        if path.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp"):
            try:
                from PIL import Image

                with Image.open(path) as image:
                    if image.width > options.image_max_width:
                        ratio = options.image_max_width / image.width
                        image = image.resize((options.image_max_width, int(image.height * ratio)))
                        image.convert("RGB").save(target, quality=88)
                        return target
            except Exception:  # noqa: BLE001 - Pillow 不可用就原样复制
                pass
        shutil.copy2(path, target)
    except OSError:
        log.exception("复制图片失败: %s", path)
        return path
    return target


# --------------------------------------------------------------------------
# 未知/二进制
# --------------------------------------------------------------------------

def read_binary_file(path: Path, options: ExtractOptions) -> ExtractedDocument:
    return ExtractedDocument(
        path=path,
        kind="binary",
        reliable=False,
        markdown=f"> 无法解析的文件类型：`{path.name}`",
        notes=["旧版 Office 格式（.doc/.ppt/.xls）无法直接解析，建议另存为 docx/pptx/xlsx"],
    )


register_reader(tuple(C.TEXT_EXTENSIONS), read_text_file)
register_reader(tuple(C.PDF_EXTENSIONS), read_pdf)
register_reader(tuple(C.DOCX_EXTENSIONS), read_docx)
register_reader(tuple(C.PPTX_EXTENSIONS), read_pptx)
register_reader(tuple(C.XLSX_EXTENSIONS), read_xlsx)
register_reader((".csv",), read_csv_file)
register_reader(tuple(C.IMAGE_EXTENSIONS), read_image_file)
register_reader(tuple(C.UNSUPPORTED_BINARY_EXTENSIONS), read_binary_file)


# --------------------------------------------------------------------------
# 统一入口
# --------------------------------------------------------------------------

def extract(path: Path, options: ExtractOptions | None = None) -> ExtractedDocument:
    """读取任意支持的文件，返回可用的 Markdown 与可靠性标记。"""
    options = options or ExtractOptions.from_config()
    path = Path(path)
    if not path.exists():
        return ExtractedDocument(path=path, reliable=False, error="文件不存在")
    if path.is_dir():
        return ExtractedDocument(path=path, reliable=False, error="这是一个文件夹")

    reader = READERS.get(path.suffix.lower())
    if reader is None:
        return ExtractedDocument(
            path=path, kind="unknown", reliable=False,
            markdown=f"> 暂不支持的文件类型：`{path.name}`",
            notes=[f"后缀 {path.suffix or '（无）'} 尚未注册读取器，可在 document_reader.READERS 中扩展"],
        )
    try:
        return reader(path, options)
    except Exception as exc:  # noqa: BLE001 - 任何解析异常都不应中断批量处理
        log.exception("读取文件失败: %s", path)
        return ExtractedDocument(path=path, reliable=False, error=f"读取失败：{exc}")


def read_text(path: Path, options: ExtractOptions | None = None) -> str:
    return extract(path, options).markdown


def supported_extensions() -> list[str]:
    return sorted(READERS.keys())


__all__ = [
    "ExtractOptions",
    "ExtractedDocument",
    "READERS",
    "extract",
    "extract_pdf_images",
    "read_text",
    "register_reader",
    "render_pdf_to_images",
    "supported_extensions",
]
